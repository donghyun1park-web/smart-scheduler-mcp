from __future__ import annotations

import json
from datetime import date

from openpyxl import load_workbook

from core import db
from core.importer import generate_weekly_report_from_db
from core.models import Activity, Calendar, CostItem, DailyRecord, Project, WBS
from core.reporting import create_weekly_construction_report


def test_weekly_report_applies_internal_hq_client_styles(tmp_path):
    data = _load_sample_data()
    outputs = {}

    for style in ("internal", "hq", "client"):
        output = tmp_path / f"{style}.xlsx"
        create_weekly_construction_report(data, output, report_style=style)
        workbook = load_workbook(output)
        report = workbook["07_보고서"]
        dashboard = workbook["05_대시보드"]
        delays = workbook["부진공정 TOP 10"]
        recovery = workbook["만회대책 초안"]
        outputs[style] = {
            "summary": report["B4"].value,
            "dashboard_opinion": dashboard["B11"].value,
            "delay": delays["H2"].value,
            "recovery": recovery["F2"].value,
        }
        workbook.close()

    assert len({outputs[style]["summary"] for style in outputs}) == 3
    assert "담당자 확인" in outputs["internal"]["summary"]
    assert "원가 집행률" in outputs["hq"]["summary"]
    assert "검토 중입니다" in outputs["client"]["summary"]
    client_text = " ".join(str(value) for value in outputs["client"].values())
    for forbidden in ["손실", "적자", "책임", "부실", "협력업체 문제", "내부 원가 과투입", "확정", "반드시"]:
        assert forbidden not in client_text


def test_db_weekly_report_passes_report_style_to_workbook(tmp_path):
    db_path = _project_db(tmp_path)
    report_path = tmp_path / "hq_from_db.xlsx"

    generate_weekly_report_from_db(db_path, report_path, report_style="hq")

    workbook = load_workbook(report_path)
    report = workbook["07_보고서"]
    assert report["B3"].value == "hq"
    assert "원가 집행률" in report["B4"].value
    workbook.close()


def _load_sample_data() -> dict[str, object]:
    with open("samples/ai_construction_site_sample.json", encoding="utf-8") as file:
        return json.load(file)


def _project_db(tmp_path):
    db_path = tmp_path / "project.scheduler"
    db.initialize_database(db_path)
    db.create_calendar(db_path, Calendar("cal-1", "Calendar", "1111100"))
    db.create_project(db_path, Project("project-1", "Demo Site", date(2026, 5, 1), "cal-1"))
    db.create_wbs(db_path, WBS("wbs-1", None, "WBS", "Root"))
    db.create_activity(
        db_path,
        Activity(
            "act-1",
            "A-100",
            "Excavation",
            "wbs-1",
            "civil",
            "B1",
            5,
            cost=1000,
        ),
    )
    db.create_daily_record(db_path, DailyRecord("dr-1", "act-1", date(2026, 5, 20), 100, 50))
    db.upsert_cost_item(db_path, CostItem("cost-1", "act-1", 1000, 900, 700, 300))
    return db_path
