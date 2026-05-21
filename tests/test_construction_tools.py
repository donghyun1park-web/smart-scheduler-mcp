from __future__ import annotations

import json
from datetime import date

from openpyxl import load_workbook

from core import db
from core.models import Calendar, ChangeLogEntry, Project
from core.recovery import FORBIDDEN_FINAL_WORDS
from tools.construction_tools import (
    detect_delays,
    generate_weekly_report,
    generate_weekly_report_from_db,
    input_daily_record,
    import_excel_input_to_db,
    list_change_log_tool,
    suggest_recovery,
    summarize_site_status,
)


def test_input_daily_record_validates_and_returns_record():
    result = input_daily_record(
        {
            "activity_id": "A-200",
            "work_date": "2026-05-21",
            "planned_qty": 100,
            "actual_qty": 65,
            "workers": 8,
            "owner": "Lee",
        }
    )

    assert result["ok"] is True
    assert result["record"]["activity_id"] == "A-200"
    assert result["validation_issues"] == []


def test_detect_delays_and_summarize_site_status_use_sample_data():
    data = _load_sample_data()

    delays = detect_delays(data, top_n=10)
    summary = summarize_site_status(data)

    assert delays["ok"] is True
    assert len(delays["delays"]) >= 4
    assert summary["ok"] is True
    assert summary["summary"]["delayed_activity_count"] >= 4


def test_suggest_recovery_returns_only_candidate_language():
    result = suggest_recovery("material_delay")
    text = json.dumps(result, ensure_ascii=False)

    assert result["ok"] is True
    assert result["plans"]
    assert "후보" in text or "초안" in text
    assert not any(word in text for word in FORBIDDEN_FINAL_WORDS)


def test_generate_weekly_report_writes_xlsx(tmp_path):
    output = tmp_path / "weekly.xlsx"

    result = generate_weekly_report(_load_sample_data(), output)

    assert result["ok"] is True
    workbook = load_workbook(result["output_path"])
    assert "05_대시보드" in workbook.sheetnames


def test_db_backed_tool_names_exist_for_mcp_surface():
    assert callable(import_excel_input_to_db)
    assert callable(generate_weekly_report_from_db)
    assert callable(list_change_log_tool)


def test_list_change_log_tool_filters_and_limits_entries(tmp_path):
    db_path = tmp_path / "project.scheduler"
    db.initialize_database(db_path)
    db.create_calendar(db_path, Calendar("cal-1", "Calendar", "1111100"))
    db.create_project(db_path, Project("project-1", "Demo", date(2026, 5, 1), "cal-1"))
    db.log_change(
        db_path,
        ChangeLogEntry(
            "chg-1",
            "daily_records",
            "act-1:2026-05-21",
            "old",
            "new",
            "excel_import_replace",
            "tester",
            changed_at=date(2026, 5, 21),
        ),
    )

    result = list_change_log_tool(str(db_path), target_table="daily_records", limit=1)

    assert result["ok"] is True
    assert result["count"] == 1
    assert result["changes"][0]["change_id"] == "chg-1"
    assert result["changes"][0]["target_table"] == "daily_records"


def _load_sample_data() -> dict[str, object]:
    with open("samples/ai_construction_site_sample.json", encoding="utf-8") as file:
        return json.load(file)
