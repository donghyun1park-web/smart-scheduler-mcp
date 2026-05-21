from __future__ import annotations

from openpyxl import load_workbook

from core.excel_io import FIELD_WORKBOOK_SHEETS, create_field_input_template


def test_create_field_input_template_has_required_sheets_and_headers(tmp_path):
    output = tmp_path / "field_template.xlsx"

    result = create_field_input_template(output, project_name="Demo Site")

    workbook = load_workbook(result["output_path"])
    assert workbook.sheetnames == FIELD_WORKBOOK_SHEETS
    assert workbook["01_실적입력"]["A1"].value == "project_name"
    assert workbook["01_실적입력"]["B1"].value == "Demo Site"
    assert workbook["02_원가입력"]["A3"].value == "activity_id"
    assert workbook["03_자재검측"]["A3"].value == "activity_id"
    assert workbook["05_대시보드"]["A1"].value == "현장소장 대시보드"


def test_template_marks_input_and_auto_cells_with_distinct_fills(tmp_path):
    output = tmp_path / "field_template.xlsx"
    create_field_input_template(output, project_name="Demo Site")

    workbook = load_workbook(output)
    input_cell = workbook["01_실적입력"]["B4"]
    auto_cell = workbook["05_대시보드"]["B3"]

    assert input_cell.fill.fgColor.rgb == "FFFFEB9C"
    assert auto_cell.fill.fgColor.rgb == "FFD9EAD3"
