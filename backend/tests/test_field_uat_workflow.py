from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from core.field_uat import FieldUATWorkflowInput, run_field_uat_workflow
from tools.field_uat_tools import run_field_uat_workflow as run_field_uat_workflow_tool


FIXTURES = Path(__file__).parent / "fixtures"


def test_field_uat_workflow_returns_summary_for_imported_schedule(tmp_path):
    imported_schedule = _fixture("hongeundong_minimal_import.json")

    result = run_field_uat_workflow(
        FieldUATWorkflowInput(
            project_id="hongeundong",
            imported_schedule=imported_schedule,
            output_dir=tmp_path,
            generate_reports=False,
        )
    )

    assert result["project_id"] == "hongeundong"
    assert result["input_summary"]["source_type"] == "imported_schedule"
    assert result["input_summary"]["task_count"] == 3
    assert result["diagnostics_summary"]["relationship_coverage_ratio"] == 0.0
    assert result["cpm_summary"]["before_finish_date"]
    assert result["recommended_next_actions"]


def test_field_uat_workflow_marks_low_relationship_coverage(tmp_path):
    imported_schedule = _fixture("hongeundong_minimal_import.json")

    result = run_field_uat_workflow(
        FieldUATWorkflowInput(
            project_id="hongeundong",
            imported_schedule=imported_schedule,
            output_dir=tmp_path,
            generate_reports=False,
        )
    )

    assert result["field_uat_status"] == "needs_relationship_correction"
    assert "Add dependencies for isolated tasks before trusting CPM finish date." in result["recommended_next_actions"]


def test_field_uat_workflow_includes_calibration_summary_when_patch_supplied(tmp_path):
    imported_schedule = _fixture("hongeundong_minimal_import.json")
    calibration_patch = _fixture("hongeundong_calibration_patch.json")

    result = run_field_uat_workflow(
        FieldUATWorkflowInput(
            project_id="hongeundong",
            imported_schedule=imported_schedule,
            calibration_patch=calibration_patch,
            target_finish_date="2026-06-12",
            output_dir=tmp_path,
            generate_reports=False,
        )
    )

    calibration_summary = result["calibration_summary"]
    assert calibration_summary is not None
    assert calibration_summary["target_finish_date"] == "2026-06-12"
    assert calibration_summary["after_finish_date"] == "2026-06-12"
    assert calibration_summary["delta_days_after"] == 0
    assert calibration_summary["correction_count"] == 3


def test_field_uat_workflow_writes_summary_json(tmp_path):
    imported_schedule = _fixture("hongeundong_minimal_import.json")

    result = run_field_uat_workflow(
        FieldUATWorkflowInput(
            project_id="hongeundong",
            imported_schedule=imported_schedule,
            output_dir=tmp_path,
            generate_reports=False,
        )
    )

    summary_path = Path(result["artifacts"]["summary_json"])
    markdown_path = Path(result["artifacts"]["summary_markdown"])
    assert summary_path.exists()
    assert markdown_path.exists()
    saved = json.loads(summary_path.read_text(encoding="utf-8"))
    assert saved["project_id"] == "hongeundong"
    assert saved["field_uat_status"] == result["field_uat_status"]


def test_run_field_uat_workflow_tool_returns_result(tmp_path):
    imported_schedule = _fixture("hongeundong_minimal_import.json")

    result = run_field_uat_workflow_tool(
        project_id="hongeundong",
        imported_schedule=imported_schedule,
        output_dir=tmp_path,
        generate_reports=False,
    )

    assert result["project_id"] == "hongeundong"
    assert result["field_uat_status"] == "needs_relationship_correction"
    assert result["artifacts"]["summary_json"]


def test_field_uat_workflow_writes_excel_summary_sheets(tmp_path):
    workbook = _field_uat_workbook(tmp_path)

    assert "Field UAT Summary" in workbook.sheetnames
    assert "Diagnostics Summary" in workbook.sheetnames
    assert "Calibration Summary" in workbook.sheetnames
    assert "Calibration Corrections" in workbook.sheetnames
    assert "Calibration Warnings" in workbook.sheetnames
    assert "Recommended Actions" in workbook.sheetnames
    workbook.close()


def test_excel_report_has_field_uat_summary_sheet(tmp_path):
    workbook = _field_uat_workbook(tmp_path)
    sheet = workbook["Field UAT Summary"]

    rows = _key_value_rows(sheet)
    assert rows["project_id"] == "hongeundong"
    assert rows["field_uat_status"] in {"needs_relationship_correction", "needs_cost_correction", "ready_for_review"}
    assert rows["summary_json_path"]
    assert rows["summary_markdown_path"]
    workbook.close()


def test_excel_report_has_diagnostics_summary_sheet(tmp_path):
    workbook = _field_uat_workbook(tmp_path)
    sheet = workbook["Diagnostics Summary"]

    rows = _key_value_rows(sheet)
    assert rows["task_count"] == 3
    assert rows["dependency_count"] == 0
    assert rows["relationship_coverage_ratio"] == 0
    workbook.close()


def test_excel_report_has_calibration_summary_sheet(tmp_path):
    workbook = _field_uat_workbook(tmp_path)
    sheet = workbook["Calibration Summary"]

    rows = _key_value_rows(sheet)
    assert rows["target_finish_date"] == "2026-06-12"
    assert rows["after_finish_date"] == "2026-06-12"
    assert rows["delta_days_after"] == 0
    workbook.close()


def test_excel_report_writes_correction_records_with_applied_order(tmp_path):
    workbook = _field_uat_workbook(tmp_path)
    sheet = workbook["Calibration Corrections"]

    headers = [cell.value for cell in sheet[1]]
    assert headers[:7] == ["applied_order", "task_id", "field", "old_value", "new_value", "reason", "source"]
    assert sheet["A2"].value == 1
    correction_fields = [row[2] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert "dependency" in correction_fields
    workbook.close()


def test_excel_report_writes_structured_warnings(tmp_path):
    workbook = _field_uat_workbook(tmp_path)
    sheet = workbook["Calibration Warnings"]

    headers = [cell.value for cell in sheet[1]]
    assert headers[:6] == ["severity", "code", "message", "task_id", "applied_order", "phase"]
    warning_codes = [row[1] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert "LOW_RELATIONSHIP_COVERAGE" in warning_codes
    workbook.close()


def _fixture(name: str) -> dict[str, object]:
    return json.loads((FIXTURES / name).read_text(encoding="utf-8"))


def _field_uat_workbook(tmp_path):
    imported_schedule = _fixture("hongeundong_minimal_import.json")
    calibration_patch = _fixture("hongeundong_calibration_patch.json")

    result = run_field_uat_workflow(
        FieldUATWorkflowInput(
            project_id="hongeundong",
            imported_schedule=imported_schedule,
            calibration_patch=calibration_patch,
            target_finish_date="2026-06-12",
            output_dir=tmp_path,
            generate_reports=True,
        )
    )

    return load_workbook(result["artifacts"]["excel_report"])


def _key_value_rows(sheet) -> dict[str, object]:
    return {
        str(row[0]): row[1]
        for row in sheet.iter_rows(min_row=1, max_col=2, values_only=True)
        if row[0] is not None
    }
