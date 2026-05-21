from __future__ import annotations

import json

from openpyxl import Workbook

from core import db
from tools import import_tools
from tools.import_tools import import_budget_excel, import_schedule_excel


def test_import_schedule_excel_dry_run_returns_json_and_does_not_create_db(tmp_path):
    excel_path = tmp_path / "schedule.xlsx"
    db_path = tmp_path / "site.scheduler"
    _write_minimal_schedule_workbook(excel_path)

    result = import_schedule_excel(str(db_path), str(excel_path))

    assert result["ok"] is True
    assert result["dry_run"] is True
    assert result["imported"]["activities"] == 1
    assert not db_path.exists()
    json.dumps(result, ensure_ascii=False)


def test_import_budget_excel_dry_run_uses_safe_temp_db(monkeypatch, tmp_path):
    excel_path = tmp_path / "budget.xlsx"
    excel_path.write_bytes(b"placeholder")
    db_path = tmp_path / "site.scheduler"
    db.initialize_database(db_path)

    calls: list[str] = []

    def fake_import_budgets_to_db(target_db, budget_paths, *, disciplines=None, project_id=None):
        calls.append(str(target_db))
        return {
            "ok": True,
            "cost_items_created": 3,
            "categories_total": 2,
            "unmatched_categories": 0,
            "warnings": [],
        }

    monkeypatch.setattr(import_tools, "import_budgets_to_db", fake_import_budgets_to_db)

    result = import_budget_excel(str(db_path), str(excel_path))

    assert result["ok"] is True
    assert result["dry_run"] is True
    assert result["imported"]["cost_items"] == 3
    assert calls and calls[0] != str(db_path)
    assert db.list_cost_items(db_path) == []
    json.dumps(result, ensure_ascii=False)


def test_import_wrappers_return_clear_missing_file_error(tmp_path):
    result = import_schedule_excel(
        str(tmp_path / "site.scheduler"),
        str(tmp_path / "missing.xlsx"),
    )

    assert result["ok"] is False
    assert result["dry_run"] is True
    assert "Excel" in str(result["error"])


def _write_minimal_schedule_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"
    ws.cell(5, 4).value = "2026"
    ws.cell(7, 4).value = 1
    ws.cell(8, 4).value = 1
    ws.cell(7, 5).value = 1
    ws.cell(8, 5).value = 11
    ws.cell(9, 1).value = "architecture"
    ws.cell(9, 2).value = "Foundation"
    ws.cell(9, 4).value = "bar"
    ws.cell(9, 5).value = "bar"
    wb.save(path)
