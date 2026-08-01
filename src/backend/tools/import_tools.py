from __future__ import annotations

import json
import shutil
import tempfile
import uuid
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from core import db
from core.backup import create_backup
from core.budget_importer import import_budgets_to_db
from core.models import Activity, Relationship, WBS
from core.schedule_importer import import_schedule_to_db
from core.validation import (
    ValidationError,
    optional_float,
    optional_text,
    require_int,
    require_text,
)


DEFAULT_PRESETS_DIR = Path("presets")
REQUIRED_COLUMNS = ("code", "name", "duration")


def import_schedule_excel(
    db_path: str,
    excel_path: str,
    *,
    project_id: str | None = None,
    dry_run: bool = True,
    backup_before_import: bool = True,
) -> dict[str, object]:
    """Import a construction schedule Excel into a .scheduler DB.

    The default is a dry run. In dry-run mode this function imports into a
    temporary DB copy and never creates or mutates the requested DB path.
    """
    db_file = Path(db_path)
    excel_file = Path(excel_path)
    if not excel_file.is_file():
        return _import_error(
            f"Excel 파일을 찾을 수 없습니다: {excel_file}",
            db_file,
            excel_file,
            dry_run=dry_run,
        )

    backup_path: str | None = None
    try:
        with _import_target(db_file, dry_run=dry_run) as target_db:
            if not dry_run and backup_before_import and db_file.is_file():
                backup = create_backup(db_file, reason="v2.3-schedule-import")
                backup_path = str(backup.backup_path)
            core_result = import_schedule_to_db(
                target_db,
                excel_file,
                project_id=project_id,
            )
    except Exception as exc:  # noqa: BLE001 - MCP tools should return structured errors.
        return _import_error(
            f"공정표 Excel import 중 오류가 발생했습니다: {exc}",
            db_file,
            excel_file,
            dry_run=dry_run,
        )

    return _normalize_import_result(
        core_result,
        db_file,
        excel_file,
        dry_run=dry_run,
        project_id=project_id or _text_or_none(core_result.get("project_id")),
        imported={
            "activities": int(core_result.get("activity_count", 0) or 0),
            "wbs": int(core_result.get("wbs_count", 0) or 0),
            "cost_items": 0,
            "relationships": int(core_result.get("relationship_count", 0) or 0),
        },
        backup_path=backup_path,
    )


def import_budget_excel(
    db_path: str,
    excel_path: str,
    *,
    project_id: str | None = None,
    dry_run: bool = True,
    backup_before_import: bool = True,
) -> dict[str, object]:
    """Import an execution-budget Excel into a .scheduler DB.

    The wrapper accepts a single Excel path for MCP ergonomics and delegates to
    the existing multi-file budget importer.
    """
    db_file = Path(db_path)
    excel_file = Path(excel_path)
    if not excel_file.is_file():
        return _import_error(
            f"Excel 파일을 찾을 수 없습니다: {excel_file}",
            db_file,
            excel_file,
            dry_run=dry_run,
        )
    if not dry_run and not db_file.is_file():
        return _import_error(
            f".scheduler DB 파일을 찾을 수 없습니다: {db_file}",
            db_file,
            excel_file,
            dry_run=dry_run,
        )

    backup_path: str | None = None
    try:
        with _import_target(db_file, dry_run=dry_run) as target_db:
            if not dry_run and backup_before_import and db_file.is_file():
                backup = create_backup(db_file, reason="v2.3-budget-import")
                backup_path = str(backup.backup_path)
            core_result = import_budgets_to_db(
                target_db,
                [excel_file],
                project_id=project_id,
            )
    except Exception as exc:  # noqa: BLE001 - MCP tools should return structured errors.
        return _import_error(
            f"실행예산 Excel import 중 오류가 발생했습니다: {exc}",
            db_file,
            excel_file,
            dry_run=dry_run,
        )

    return _normalize_import_result(
        core_result,
        db_file,
        excel_file,
        dry_run=dry_run,
        project_id=project_id,
        imported={
            "activities": 0,
            "wbs": 0,
            "cost_items": int(core_result.get("cost_items_created", 0) or 0),
            "relationships": 0,
        },
        backup_path=backup_path,
    )


def save_column_mapping_preset(
    preset_name: str,
    columns: dict[str, str],
    presets_dir: str | Path = DEFAULT_PRESETS_DIR,
) -> dict[str, object]:
    root = Path(presets_dir)
    root.mkdir(parents=True, exist_ok=True)
    path = root / f"{preset_name}.json"
    path.write_text(
        json.dumps({"name": preset_name, "columns": columns}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return {"ok": True, "preset_path": str(path)}


def load_column_mapping_preset(
    preset_name: str,
    presets_dir: str | Path = DEFAULT_PRESETS_DIR,
) -> dict[str, Any]:
    path = Path(presets_dir) / f"{preset_name}.json"
    return json.loads(path.read_text(encoding="utf-8"))


class _ImportTarget:
    def __init__(self, db_path: Path, *, dry_run: bool) -> None:
        self.db_path = db_path
        self.dry_run = dry_run
        self._tmpdir: tempfile.TemporaryDirectory[str] | None = None
        self.target_path = db_path

    def __enter__(self) -> Path:
        if not self.dry_run:
            return self.db_path
        self._tmpdir = tempfile.TemporaryDirectory(
            prefix="smart_scheduler_import_",
            ignore_cleanup_errors=True,
        )
        target_name = self.db_path.name or "dry_run.scheduler"
        self.target_path = Path(self._tmpdir.name) / target_name
        if self.db_path.is_file():
            shutil.copy2(self.db_path, self.target_path)
        return self.target_path

    def __exit__(self, exc_type, exc, tb) -> None:
        if self._tmpdir is not None:
            self._tmpdir.cleanup()


def _import_target(db_path: Path, *, dry_run: bool) -> _ImportTarget:
    return _ImportTarget(db_path, dry_run=dry_run)


def _import_error(
    message: str,
    db_path: Path,
    excel_path: Path,
    *,
    dry_run: bool,
) -> dict[str, object]:
    return {
        "ok": False,
        "dry_run": dry_run,
        "db_path": str(db_path),
        "excel_path": str(excel_path),
        "error": message,
        "warnings": [],
        "errors": [message],
    }


def _normalize_import_result(
    core_result: dict[str, Any],
    db_path: Path,
    excel_path: Path,
    *,
    dry_run: bool,
    project_id: str | None,
    imported: dict[str, int],
    backup_path: str | None,
) -> dict[str, object]:
    warnings = [str(item) for item in core_result.get("warnings", [])]
    errors = [str(item) for item in core_result.get("errors", [])]
    ok = bool(core_result.get("ok", not errors))
    result: dict[str, object] = {
        "ok": ok,
        "dry_run": dry_run,
        "db_path": str(db_path),
        "excel_path": str(excel_path),
        "project_id": project_id,
        "imported": imported,
        "created_ids": list(core_result.get("created_ids", [])),
        "updated_ids": list(core_result.get("updated_ids", [])),
        "warnings": warnings,
        "errors": errors,
    }
    if backup_path:
        result["backup_path"] = backup_path
    if not ok and errors:
        result["error"] = errors[0]
    elif not ok:
        result["error"] = str(core_result.get("error") or "Excel 가져오기에 실패했습니다.")
    return result


def _text_or_none(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def import_excel(
    project_path: str | Path,
    file_path: str | Path,
    preset_name: str | None = None,
    header_row: int = 1,
    presets_dir: str | Path = DEFAULT_PRESETS_DIR,
    column_mapping: dict[str, str] | None = None,
) -> dict[str, object]:
    """Import activities (and predecessor relationships) from an Excel workbook.

    If neither ``preset_name`` nor ``column_mapping`` is supplied, the importer
    falls back to the Korean timeline layout (header row 5, data from row 9).

    The mapping must cover ``code``, ``name``, ``duration``, and one of
    ``wbs_code``/``wbs_id``. Optional keys: ``discipline``, ``zone``, ``cost``,
    ``predecessors`` (comma-separated activity codes resolved after all rows
    are read, so forward references are allowed).

    Returns row-level ``failed_rows`` instead of raising — callers can surface
    each reason to the user. ``ok`` is True only when every row imported.
    """
    if column_mapping is None and preset_name is None:
        return _import_korean_timeline_schedule(project_path, file_path)

    mapping = column_mapping or _load_mapping(preset_name, presets_dir)
    _validate_required_mapping(mapping)
    warnings: list[str] = []
    failed_rows: list[dict[str, object]] = []
    added_activities = 0
    added_relationships = 0
    code_to_activity_id: dict[str, str] = {}

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row))]
        header_index = {str(value).strip(): idx for idx, value in enumerate(headers) if value is not None}
        missing_headers = [label for label in mapping.values() if label not in header_index]
        if missing_headers:
            return {
                "ok": False,
                "error_code": "MISSING_HEADERS",
                "error_message": f"Missing mapped headers: {missing_headers}",
                "warnings": [f"Missing mapped headers: {missing_headers}"],
                "failed_rows": [],
                "added_activities": 0,
                "added_relationships": 0,
            }
        wbs_by_code = {wbs.code: wbs for wbs in db.list_wbs(project_path)}
        pending_predecessors: list[tuple[int, str, str]] = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            try:
                values = _row_values(row, header_index, mapping)
                code = require_text(values.get("code"), "code")
                if code in code_to_activity_id:
                    raise ValidationError(f"Duplicate activity code: {code}")
                wbs_code = optional_text(
                    values.get("wbs_code") or values.get("wbs_id"),
                    default="ROOT",
                )
                wbs_id = _ensure_wbs(project_path, wbs_by_code, wbs_code)
                duration = require_int(values.get("duration"), "duration", minimum=0)
                cost = optional_float(values.get("cost"), "cost", default=0.0)
                activity_id = str(uuid.uuid4())
                db.add_activity(
                    project_path,
                    Activity(
                        activity_id=activity_id,
                        code=code,
                        name=require_text(values.get("name"), "name"),
                        wbs_id=wbs_id,
                        discipline=optional_text(values.get("discipline"), default="공통"),
                        zone=optional_text(values.get("zone")),
                        duration=duration,
                        cost=cost,
                    ),
                )
                code_to_activity_id[code] = activity_id
                added_activities += 1
                predecessors = optional_text(values.get("predecessors"))
                if predecessors:
                    pending_predecessors.append((row_number, activity_id, predecessors))
            except ValidationError as exc:
                failed_rows.append({"row": row_number, "reason": str(exc)})
            except Exception as exc:  # noqa: BLE001 - unexpected row-level errors must be reported.
                failed_rows.append({"row": row_number, "reason": str(exc)})

        for row_number, succ_id, predecessors in pending_predecessors:
            for pred_code in [item.strip() for item in predecessors.split(",") if item.strip()]:
                pred_id = code_to_activity_id.get(pred_code)
                if pred_id is None:
                    failed_rows.append({"row": row_number, "reason": f"Unknown predecessor: {pred_code}"})
                    continue
                db.add_relationship(
                    project_path,
                    Relationship(str(uuid.uuid4()), pred_id, succ_id, "FS", 0),
                )
                added_relationships += 1
    finally:
        workbook.close()

    return {
        "ok": len(failed_rows) == 0,
        "warnings": warnings,
        "failed_rows": failed_rows,
        "added_activities": added_activities,
        "added_relationships": added_relationships,
    }


def _import_korean_timeline_schedule(project_path: str | Path, file_path: str | Path) -> dict[str, object]:
    warnings = [
        "Detected Korean timeline schedule format; imported row-level activities with inferred durations.",
        "No predecessor/relationship column was found. CPM can run, but Critical Path and completion date require manual relationship correction.",
        "No cost column was found. S-Curve will be empty until costs are mapped or entered manually.",
    ]
    failed_rows: list[dict[str, object]] = []
    added_activities = 0
    wbs_by_code = {wbs.code: wbs for wbs in db.list_wbs(project_path)}

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if not _looks_like_korean_timeline(sheet):
            raise ValueError("preset_name or column_mapping is required")

        current_wbs_code = "ROOT"
        for row_number, row in enumerate(sheet.iter_rows(min_row=9, values_only=True), start=9):
            discipline_cell = _clean(row[0] if len(row) > 0 else None)
            item_cell = _clean(row[1] if len(row) > 1 else None)
            timeline_values = [_clean(value) for value in row[3:]]
            timeline_markers = [(idx, value) for idx, value in enumerate(timeline_values) if value]

            if discipline_cell and discipline_cell != "주요 Milestone":
                current_wbs_code = discipline_cell
                _ensure_wbs(project_path, wbs_by_code, current_wbs_code)

            if not timeline_markers or discipline_cell == "주요 Milestone":
                continue
            try:
                activity_id = str(uuid.uuid4())
                code = f"G{row_number:03d}"
                name = item_cell or timeline_markers[0][1]
                duration = _infer_timeline_duration(timeline_markers)
                db.add_activity(
                    project_path,
                    Activity(
                        activity_id=activity_id,
                        code=code,
                        name=name,
                        wbs_id=_ensure_wbs(project_path, wbs_by_code, current_wbs_code),
                        discipline=_infer_discipline(current_wbs_code),
                        zone="",
                        duration=duration,
                        cost=0.0,
                    ),
                )
                added_activities += 1
            except Exception as exc:  # noqa: BLE001 - row-level import errors must be reported.
                failed_rows.append({"row": row_number, "reason": str(exc)})
    finally:
        workbook.close()

    return {
        "ok": len(failed_rows) == 0,
        "import_mode": "korean_timeline",
        "warnings": warnings,
        "failed_rows": failed_rows,
        "added_activities": added_activities,
        "added_relationships": 0,
    }


def _looks_like_korean_timeline(sheet: Any) -> bool:
    first_headers = [sheet.cell(row=5, column=col).value for col in range(1, 4)]
    return [_normalize_header(value) for value in first_headers] == ["공종", "항목", "구분"]


def _clean(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _normalize_header(value: object) -> str:
    return _clean(value).replace(" ", "")


def _infer_timeline_duration(markers: list[tuple[int, str]]) -> int:
    first = markers[0][0]
    last = markers[-1][0]
    return max(1, (last - first + 1) * 10)


def _infer_discipline(wbs_code: str) -> str:
    if "위생" in wbs_code:
        return "위생"
    if "공조" in wbs_code or "기계" in wbs_code:
        return "공조"
    if "소방" in wbs_code:
        return "소방"
    if "전기" in wbs_code:
        return "전기"
    if "제어" in wbs_code:
        return "자동제어"
    return "공통"


def _load_mapping(preset_name: str | None, presets_dir: str | Path) -> dict[str, str]:
    if preset_name is None:
        raise ValueError("preset_name or column_mapping is required")
    preset = load_column_mapping_preset(preset_name, presets_dir)
    return dict(preset["columns"])


def _validate_required_mapping(mapping: dict[str, str]) -> None:
    missing = [column for column in REQUIRED_COLUMNS if column not in mapping]
    if "wbs_code" not in mapping and "wbs_id" not in mapping:
        missing.append("wbs_code or wbs_id")
    if missing:
        raise ValueError(f"Missing required column mapping: {missing}")


def _row_values(row: tuple[object, ...], header_index: dict[str, int], mapping: dict[str, str]) -> dict[str, object]:
    values: dict[str, object] = {}
    for standard_name, header_name in mapping.items():
        values[standard_name] = row[header_index[header_name]]
    return values


def _ensure_wbs(project_path: str | Path, wbs_by_code: dict[str, WBS], code: str) -> str:
    existing = wbs_by_code.get(code)
    if existing is not None:
        return existing.wbs_id
    wbs = db.add_wbs(project_path, WBS(str(uuid.uuid4()), None, code, code, len(wbs_by_code) + 1))
    wbs_by_code[code] = wbs
    return wbs.wbs_id
