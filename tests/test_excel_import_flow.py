from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from core import db
from core.excel_io import create_field_input_template
from core.importer import generate_weekly_report_from_db, import_field_input_to_db
from core.models import Activity, Calendar, Project, WBS


def test_excel_input_imports_to_db_and_generates_weekly_report(tmp_path):
    db_path = _project_db(tmp_path)
    input_path = tmp_path / "현장입력.xlsx"
    report_path = tmp_path / "weekly_from_db.xlsx"
    create_field_input_template(
        input_path,
        project_name="Demo Site",
        activities=[{"activity_id": "act-1"}],
    )
    workbook = load_workbook(input_path)
    workbook["01_실적입력"]["B4"] = "2026-05-21"
    workbook["01_실적입력"]["C4"] = 100
    workbook["01_실적입력"]["D4"] = 70
    workbook["01_실적입력"]["E4"] = 6
    workbook["01_실적입력"]["G4"] = "Kim"
    workbook["02_원가입력"]["B4"] = 1000
    workbook["02_원가입력"]["C4"] = 900
    workbook["02_원가입력"]["D4"] = 700
    workbook["02_원가입력"]["E4"] = 300
    workbook["03_자재검측"]["B4"] = "rebar"
    workbook["03_자재검측"]["C4"] = "2026-05-20"
    workbook["03_자재검측"]["H4"] = "ordered"
    workbook.save(input_path)
    workbook.close()

    import_result = import_field_input_to_db(db_path, input_path, project_id="project-1", imported_by="tester")
    report_result = generate_weekly_report_from_db(db_path, report_path)

    assert import_result["daily_records_inserted"] == 1
    assert import_result["cost_items_upserted"] == 1
    assert import_result["materials_inserted"] == 1
    assert db.get_cumulative_qty(db_path, "act-1")["progress_pct"] == 70.0
    assert db.get_cost_summary(db_path)["cost_execution_rate"] == 77.78
    assert report_result["ok"] is True
    workbook = load_workbook(report_result["output_path"])
    assert {
        "05_대시보드",
        "금주 실적",
        "부진공정 TOP 10",
        "만회대책 초안",
        "원가현황",
        "다음 주 예정공정",
        "간트 데이터",
        "07_보고서",
    }.issubset(set(workbook.sheetnames))
    workbook.close()


def _project_db(tmp_path):
    db_path = tmp_path / "project.scheduler"
    db.initialize_database(db_path)
    db.create_calendar(db_path, Calendar("cal-1", "Calendar", "1111100"))
    db.create_project(db_path, Project("project-1", "Demo Site", date(2026, 5, 1), "cal-1"))
    db.create_wbs(db_path, WBS("wbs-1", None, "WBS", "Root"))
    db.create_activity(
        db_path,
        Activity(
            "act-1",
            "A-100",
            "Rebar assembly",
            "wbs-1",
            "civil",
            "B1",
            5,
            cost=1000,
            es_date=date(2026, 5, 20),
            ef_date=date(2026, 5, 24),
        ),
    )
    return db_path
