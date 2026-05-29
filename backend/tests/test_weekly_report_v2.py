from __future__ import annotations

import json

from openpyxl import load_workbook

from core.reporting import create_weekly_construction_report


def test_create_weekly_construction_report_contains_required_v2_sheets(tmp_path):
    data = _load_sample_data()
    output = tmp_path / "weekly_report.xlsx"

    result = create_weekly_construction_report(data, output)

    workbook = load_workbook(result["output_path"])
    assert {
        "05_대시보드",
        "금주 실적",
        "부진공정 TOP 10",
        "만회대책 초안",
        "원가현황",
        "다음 주 예정공정",
        "간트 데이터",
    }.issubset(set(workbook.sheetnames))
    assert result["delayed_count"] >= 4
    assert result["cost_overrun_count"] >= 1


def test_weekly_report_writes_dashboard_delay_cost_and_recovery_content(tmp_path):
    data = _load_sample_data()
    output = tmp_path / "weekly_report.xlsx"

    create_weekly_construction_report(data, output, report_style="owner")

    workbook = load_workbook(output)
    dashboard = workbook["05_대시보드"]
    delays = workbook["부진공정 TOP 10"]
    recovery = workbook["만회대책 초안"]
    costs = workbook["원가현황"]
    gantt = workbook["간트 데이터"]

    assert dashboard["A1"].value == "현장소장 대시보드"
    assert dashboard["A3"].value == "전체 실적공정률"
    assert delays["A2"].value == "A-200"
    assert "후보" in recovery["D2"].value
    assert "검토 필요" in recovery["E2"].value
    assert costs["A2"].value == "A-300"
    assert gantt["A1"].value == "activity_id"


def _load_sample_data() -> dict[str, object]:
    with open("samples/ai_construction_site_sample.json", encoding="utf-8") as file:
        return json.load(file)
