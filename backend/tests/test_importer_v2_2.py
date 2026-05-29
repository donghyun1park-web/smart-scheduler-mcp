from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from core import db
from core.excel_io import create_field_input_template
from core.importer import import_field_input_to_db
from core.models import Activity, Calendar, DailyRecord, Project, WBS


def test_import_fail_policy_blocks_duplicate_without_changes_or_backup(tmp_path):
    db_path = _project_db(tmp_path)
    db.create_daily_record(db_path, DailyRecord("existing", "act-1", date(2026, 5, 21), 100, 20))
    input_path = _daily_input(tmp_path, planned_qty=100, actual_qty=80)

    result = import_field_input_to_db(db_path, input_path, conflict_policy="fail", actor="tester")

    assert result["ok"] is False
    assert result["changed"] is False
    assert result["backup_path"] is None
    assert result["conflicts"]
    assert db.get_cumulative_qty(db_path, "act-1")["actual_qty"] == 20.0


def test_import_skip_policy_keeps_existing_duplicate_and_records_skip(tmp_path):
    db_path = _project_db(tmp_path)
    db.create_daily_record(db_path, DailyRecord("existing", "act-1", date(2026, 5, 21), 100, 20))
    input_path = _daily_input(tmp_path, planned_qty=100, actual_qty=80)

    result = import_field_input_to_db(db_path, input_path, conflict_policy="skip", actor="tester")

    assert result["ok"] is True
    assert result["changed"] is True
    assert result["skipped"]["daily_records"] == 1
    assert result["backup_path"]
    assert db.get_cumulative_qty(db_path, "act-1")["actual_qty"] == 20.0


def test_import_replace_policy_replaces_duplicate_and_logs_change(tmp_path):
    db_path = _project_db(tmp_path)
    db.create_daily_record(db_path, DailyRecord("existing", "act-1", date(2026, 5, 21), 100, 20))
    input_path = _daily_input(tmp_path, planned_qty=100, actual_qty=80)

    result = import_field_input_to_db(db_path, input_path, conflict_policy="replace", actor="tester")

    assert result["ok"] is True
    assert result["updated"]["daily_records"] == 1
    assert result["backup_path"]
    assert db.get_cumulative_qty(db_path, "act-1")["actual_qty"] == 80.0
    changes = db.list_change_log(db_path, target_table="daily_records")
    assert any(change.reason == "excel_import_replace" for change in changes)


def test_import_dry_run_validates_without_backup_or_db_changes(tmp_path):
    db_path = _project_db(tmp_path)
    input_path = _daily_input(tmp_path, planned_qty=100, actual_qty=80)

    result = import_field_input_to_db(db_path, input_path, dry_run=True)

    assert result["ok"] is True
    assert result["dry_run"] is True
    assert result["changed"] is False
    assert result["backup_path"] is None
    assert db.list_daily_records(db_path) == []


def test_import_batch_validation_blocks_bad_daily_rows(tmp_path):
    db_path = _project_db(tmp_path)
    input_path = _daily_input(tmp_path, activity_id="act-1", planned_qty=-1, actual_qty=5)

    result = import_field_input_to_db(db_path, input_path)

    assert result["ok"] is False
    assert result["changed"] is False
    assert any("planned_qty" in error and "must be >= 0" in error for error in result["errors"])
    assert db.list_daily_records(db_path) == []


def test_import_batch_validation_blocks_unknown_activity(tmp_path):
    db_path = _project_db(tmp_path)
    input_path = _daily_input(tmp_path, activity_id="missing", planned_qty=10, actual_qty=5)

    result = import_field_input_to_db(db_path, input_path)

    assert result["ok"] is False
    assert any("unknown activity_id" in error for error in result["errors"])


def test_import_cost_progress_gap_returns_warning(tmp_path):
    db_path = _project_db(tmp_path)
    input_path = _daily_input(tmp_path, planned_qty=100, actual_qty=10)
    workbook = load_workbook(input_path)
    cost = workbook["02_원가입력"]
    cost["B4"] = 1000
    cost["C4"] = 1000
    cost["D4"] = 800
    cost["E4"] = 100
    workbook.save(input_path)
    workbook.close()

    result = import_field_input_to_db(db_path, input_path, dry_run=True)

    assert result["ok"] is True
    assert any("cost overrun risk" in warning for warning in result["warnings"])


def test_import_merge_policy_is_explicitly_unsupported(tmp_path):
    db_path = _project_db(tmp_path)
    input_path = _daily_input(tmp_path, planned_qty=100, actual_qty=80)

    result = import_field_input_to_db(db_path, input_path, conflict_policy="merge")

    assert result["ok"] is False
    assert any("not implemented" in error for error in result["errors"])


def _project_db(tmp_path):
    db_path = tmp_path / "project.scheduler"
    db.initialize_database(db_path)
    db.create_calendar(db_path, Calendar("cal-1", "Calendar", "1111100"))
    db.create_project(db_path, Project("project-1", "Demo", date(2026, 5, 1), "cal-1"))
    db.create_wbs(db_path, WBS("wbs-1", None, "WBS", "Root"))
    db.create_activity(db_path, Activity("act-1", "A-100", "Activity", "wbs-1", "civil", "B1", 5))
    return db_path


def _daily_input(tmp_path, *, activity_id: str = "act-1", planned_qty: float, actual_qty: float):
    input_path = tmp_path / "field_input.xlsx"
    create_field_input_template(input_path, project_name="Demo", activities=[{"activity_id": activity_id}])
    workbook = load_workbook(input_path)
    daily = workbook["01_실적입력"]
    daily["B4"] = "2026-05-21"
    daily["C4"] = planned_qty
    daily["D4"] = actual_qty
    daily["E4"] = 3
    daily["G4"] = "Kim"
    workbook.save(input_path)
    workbook.close()
    return input_path
