from __future__ import annotations

import json
import uuid
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook

from core import db
from core.models import Activity, Relationship, WBS


DEFAULT_PRESETS_DIR = Path("presets")
REQUIRED_COLUMNS = ("code", "name", "duration")


def save_column_mapping_preset(
    preset_name: str,
    columns: dict[str, str],
    presets_dir: str | Path = DEFAULT_PRESETS_DIR,
) -> dict[str, object]:
    root = Path(presets_dir)
    root.mkdir(parents=True, exist_ok=True)
    path = root / f"{preset_name}.json"
    path.write_text(
        json.dumps({"name": preset_name, "columns": columns}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return {"ok": True, "preset_path": str(path)}


def load_column_mapping_preset(
    preset_name: str,
    presets_dir: str | Path = DEFAULT_PRESETS_DIR,
) -> dict[str, Any]:
    path = Path(presets_dir) / f"{preset_name}.json"
    return json.loads(path.read_text(encoding="utf-8"))


def import_excel(
    project_path: str | Path,
    file_path: str | Path,
    preset_name: str | None = None,
    header_row: int = 1,
    presets_dir: str | Path = DEFAULT_PRESETS_DIR,
    column_mapping: dict[str, str] | None = None,
) -> dict[str, object]:
    mapping = column_mapping or _load_mapping(preset_name, presets_dir)
    _validate_required_mapping(mapping)
    warnings: list[str] = []
    failed_rows: list[dict[str, object]] = []
    added_activities = 0
    added_relationships = 0
    code_to_activity_id: dict[str, str] = {}

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        headers = [cell.value for cell in next(sheet.iter_rows(min_row=header_row, max_row=header_row))]
        header_index = {str(value).strip(): idx for idx, value in enumerate(headers) if value is not None}
        missing_headers = [label for label in mapping.values() if label not in header_index]
        if missing_headers:
            return {
                "ok": False,
                "warnings": [f"Missing mapped headers: {missing_headers}"],
                "failed_rows": [],
                "added_activities": 0,
                "added_relationships": 0,
            }
        wbs_by_code = {wbs.code: wbs for wbs in db.list_wbs(project_path)}
        pending_predecessors: list[tuple[int, str, str]] = []
        for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            try:
                values = _row_values(row, header_index, mapping)
                code = str(values["code"]).strip()
                wbs_id = _ensure_wbs(project_path, wbs_by_code, str(values.get("wbs_code") or values.get("wbs_id") or "ROOT"))
                duration = int(cast(str | int | float, values["duration"]))
                cost = float(cast(str | int | float, values.get("cost") or 0))
                activity_id = str(uuid.uuid4())
                db.add_activity(
                    project_path,
                    Activity(
                        activity_id=activity_id,
                        code=code,
                        name=str(values["name"]).strip(),
                        wbs_id=wbs_id,
                        discipline=str(values.get("discipline") or "공통").strip(),
                        zone=str(values.get("zone") or "").strip(),
                        duration=duration,
                        cost=cost,
                    ),
                )
                code_to_activity_id[code] = activity_id
                added_activities += 1
                predecessors = str(values.get("predecessors") or "").strip()
                if predecessors:
                    pending_predecessors.append((row_number, activity_id, predecessors))
            except Exception as exc:  # noqa: BLE001 - row-level import errors must be reported.
                failed_rows.append({"row": row_number, "reason": str(exc)})

        for row_number, succ_id, predecessors in pending_predecessors:
            for pred_code in [item.strip() for item in predecessors.split(",") if item.strip()]:
                pred_id = code_to_activity_id.get(pred_code)
                if pred_id is None:
                    failed_rows.append({"row": row_number, "reason": f"Unknown predecessor: {pred_code}"})
                    continue
                db.add_relationship(
                    project_path,
                    Relationship(str(uuid.uuid4()), pred_id, succ_id, "FS", 0),
                )
                added_relationships += 1
    finally:
        workbook.close()

    return {
        "ok": len(failed_rows) == 0,
        "warnings": warnings,
        "failed_rows": failed_rows,
        "added_activities": added_activities,
        "added_relationships": added_relationships,
    }


def _load_mapping(preset_name: str | None, presets_dir: str | Path) -> dict[str, str]:
    if preset_name is None:
        raise ValueError("preset_name or column_mapping is required")
    preset = load_column_mapping_preset(preset_name, presets_dir)
    return dict(preset["columns"])


def _validate_required_mapping(mapping: dict[str, str]) -> None:
    missing = [column for column in REQUIRED_COLUMNS if column not in mapping]
    if "wbs_code" not in mapping and "wbs_id" not in mapping:
        missing.append("wbs_code or wbs_id")
    if missing:
        raise ValueError(f"Missing required column mapping: {missing}")


def _row_values(row: tuple[object, ...], header_index: dict[str, int], mapping: dict[str, str]) -> dict[str, object]:
    values: dict[str, object] = {}
    for standard_name, header_name in mapping.items():
        values[standard_name] = row[header_index[header_name]]
    return values


def _ensure_wbs(project_path: str | Path, wbs_by_code: dict[str, WBS], code: str) -> str:
    existing = wbs_by_code.get(code)
    if existing is not None:
        return existing.wbs_id
    wbs = db.add_wbs(project_path, WBS(str(uuid.uuid4()), None, code, code, len(wbs_by_code) + 1))
    wbs_by_code[code] = wbs
    return wbs.wbs_id
