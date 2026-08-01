"""Import construction schedule bar-chart Excel into smart-scheduler DB.

Supports the common Korean construction schedule format where:
- Rows 5-8 form the header (year / month-group / month / day-of-month)
- Col A = 공종 (discipline), Col B = 항목 (item), Col C = 구분 (detail)
- Cols D onwards contain bar-chart cells (10-day periods per column)

The importer reads the header to build a column→date mapping, then scans
activity rows to detect bar spans (first/last non-empty cell in schedule
columns) and converts them to Activity records with es_date / ef_date.
"""
from __future__ import annotations

import re
import uuid
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


# ---------------------------------------------------------------------------
# Public dataclass for parsed activities
# ---------------------------------------------------------------------------

@dataclass
class ParsedActivity:
    """A single activity extracted from a bar-chart schedule Excel."""

    row_number: int
    discipline: str
    item: str
    detail: str
    es_date: date | None
    ef_date: date | None
    duration_days: int
    bar_texts: list[str]


@dataclass
class ScheduleParseResult:
    """Result of parsing a construction schedule workbook."""

    project_name: str
    project_start: date | None
    project_end: date | None
    scale: str
    activities: list[ParsedActivity]
    warnings: list[str]
    sheet_name: str
    column_date_map: dict[int, date]


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def parse_schedule_excel(
    path: str | Path,
    *,
    sheet_name: str | None = None,
    header_row: int = 5,
    first_data_row: int = 9,
    schedule_start_col: int = 4,
) -> ScheduleParseResult:
    """Parse a Korean bar-chart schedule Excel and return structured activities.

    Parameters
    ----------
    path : file path to the .xlsx workbook
    sheet_name : which sheet to read (default: first sheet)
    header_row : row number where '공 종' / '항 목' / '구 분' appears (1-based)
    first_data_row : first row that can contain activity data (1-based)
    schedule_start_col : first column of the schedule area (1-based, default=4=col D)
    """
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    warnings: list[str] = []

    # --- Parse project metadata from top rows ---
    project_name, project_start, project_end, scale = _parse_project_header(ws)

    # --- Build column → date mapping ---
    col_date_map = _build_column_date_map(
        ws,
        header_row=header_row,
        schedule_start_col=schedule_start_col,
        project_start=project_start,
    )
    if not col_date_map:
        warnings.append("헤더 행에서 날짜 열 매핑을 만들지 못했습니다. 헤더 행 번호를 확인하세요.")

    # --- Scan activity rows ---
    raw_activities = _scan_activities(
        ws,
        col_date_map=col_date_map,
        first_data_row=first_data_row,
        schedule_start_col=schedule_start_col,
        warnings=warnings,
    )

    # --- Merge nameless continuation rows into previous activity ---
    activities = _merge_continuation_rows(raw_activities)

    # --- Fix missing ef_date using project_end or max known date ---
    max_date = max(col_date_map.values()) if col_date_map else project_end
    for a in activities:
        if a.es_date and not a.ef_date and max_date:
            a.ef_date = max_date
            a.duration_days = max((max_date - a.es_date).days, 1)

    wb.close()
    return ScheduleParseResult(
        project_name=project_name,
        project_start=project_start,
        project_end=project_end,
        scale=scale,
        activities=activities,
        warnings=warnings,
        sheet_name=ws.title,
        column_date_map=col_date_map,
    )


def import_schedule_to_db(
    db_path: str | Path,
    schedule_path: str | Path,
    *,
    sheet_name: str | None = None,
    calendar_id: str = "cal-default",
    project_id: str | None = None,
) -> dict[str, Any]:
    """Parse a schedule Excel and write Project + WBS + Activities into a .scheduler DB.

    Returns a summary dict with ok, project_id, activity_count, warnings.
    """
    from core import db
    from core.models import Activity, Calendar, Project, WBS

    parsed = parse_schedule_excel(schedule_path, sheet_name=sheet_name)
    errors: list[str] = []

    if not parsed.activities:
        errors.append("공정표 Excel에서 활동(Activity)을 찾지 못했습니다.")
    if errors:
        return {"ok": False, "errors": errors, "warnings": parsed.warnings}

    pid = project_id or f"proj-{uuid.uuid4().hex[:8]}"
    start = parsed.project_start or date.today()

    db.initialize_database(db_path)

    # Calendar
    try:
        db.create_calendar(db_path, Calendar(calendar_id, "기본달력", "1111100"))
    except Exception:
        pass  # already exists

    # Project
    try:
        db.create_project(db_path, Project(pid, parsed.project_name or "Imported Project", start, calendar_id))
    except Exception:
        pass  # already exists

    # WBS — group by discipline
    disciplines = sorted({a.discipline for a in parsed.activities if a.discipline})
    wbs_root_id = f"wbs-{pid}-root"
    try:
        db.create_wbs(db_path, WBS(wbs_root_id, None, "ROOT", parsed.project_name or "Root"))
    except Exception:
        pass

    discipline_wbs: dict[str, str] = {}
    for idx, disc in enumerate(disciplines):
        wbs_id = f"wbs-{pid}-{idx:03d}"
        try:
            db.create_wbs(db_path, WBS(wbs_id, wbs_root_id, disc, disc, sort_order=idx))
        except Exception:
            pass
        discipline_wbs[disc] = wbs_id

    # Activities
    created = 0
    for idx, pa in enumerate(parsed.activities):
        wbs_id = discipline_wbs.get(pa.discipline, wbs_root_id)
        activity_id = f"act-{pid}-{idx:03d}"
        code = f"A-{idx + 1:03d}"
        name = pa.item or pa.detail or pa.discipline or f"Activity {idx + 1}"
        if pa.detail and pa.item and pa.detail != pa.item:
            name = f"{pa.item} - {pa.detail}"

        discipline = _normalize_discipline(pa.discipline)
        zone = ""

        try:
            db.create_activity(
                db_path,
                Activity(
                    activity_id=activity_id,
                    code=code,
                    name=name,
                    wbs_id=wbs_id,
                    discipline=discipline,
                    zone=zone,
                    duration=max(pa.duration_days, 1),
                    es_date=pa.es_date,
                    ef_date=pa.ef_date,
                ),
            )
            created += 1
        except Exception as exc:
            parsed.warnings.append(f"{pa.row_number}행 '{name}' 활동 생성 실패: {exc}")

    return {
        "ok": True,
        "project_id": pid,
        "project_name": parsed.project_name,
        "activity_count": created,
        "total_parsed": len(parsed.activities),
        "warnings": parsed.warnings,
        "db_path": str(db_path),
    }


# ---------------------------------------------------------------------------
# Internal: header parsing
# ---------------------------------------------------------------------------

_DATE_PATTERN = re.compile(r"(\d{4})\s*년\s*(\d{1,2})\s*월\s*(\d{1,2})\s*일")
_YEAR_PATTERN = re.compile(r"(\d{4})\s*년")


def _parse_project_header(ws: Any) -> tuple[str, date | None, date | None, str]:
    """Extract project name, start/end dates, and scale from the top rows."""
    project_name = ""
    project_start: date | None = None
    project_end: date | None = None
    scale = ""

    for r in range(1, 8):
        cell_val = str(ws.cell(r, 1).value or "")
        normalized = cell_val.replace(" ", "")
        if "공사명" in normalized:
            # Extract name after colon
            parts = cell_val.split(":", 1)
            if len(parts) > 1:
                project_name = parts[1].strip()
            else:
                project_name = cell_val.strip()
        elif "공사기간" in normalized:
            dates = _DATE_PATTERN.findall(cell_val)
            if dates:
                try:
                    project_start = date(int(dates[0][0]), int(dates[0][1]), int(dates[0][2]))
                except ValueError:
                    pass
            if len(dates) > 1:
                try:
                    project_end = date(int(dates[1][0]), int(dates[1][1]), int(dates[1][2]))
                except ValueError:
                    pass
        elif "공사규모" in normalized:
            parts = cell_val.split(":", 1)
            if len(parts) > 1:
                scale = parts[1].strip()
            else:
                scale = cell_val

    return project_name, project_start, project_end, scale


def _build_column_date_map(
    ws: Any,
    *,
    header_row: int,
    schedule_start_col: int,
    project_start: date | None,
) -> dict[int, date]:
    """Build a mapping from column number (1-based) to date.

    Uses rows: header_row (year), header_row+1 (month-group), header_row+2 (month), header_row+3 (day).
    """
    year_row = header_row
    month_group_row = header_row + 1
    month_row = header_row + 2
    day_row = header_row + 3

    max_col = ws.max_column or 200

    # Step 1: read year markers and forward-fill
    col_year: dict[int, int] = {}
    current_year: int | None = None
    for c in range(schedule_start_col, max_col + 1):
        val = str(ws.cell(year_row, c).value or "")
        match = _YEAR_PATTERN.search(val)
        if match:
            current_year = int(match.group(1))
        if current_year is not None:
            col_year[c] = current_year

    # Step 2: read month and day for each column
    col_date_map: dict[int, date] = {}
    for c in range(schedule_start_col, max_col + 1):
        year = col_year.get(c)
        month_val = ws.cell(month_row, c).value
        day_val = ws.cell(day_row, c).value
        if year is None or month_val is None or day_val is None:
            continue
        try:
            m = int(month_val)
            d = int(day_val)
            if 1 <= m <= 12 and 1 <= d <= 31:
                col_date_map[c] = date(year, m, min(d, 28))
        except (ValueError, TypeError):
            continue

    # Step 3: handle "착공-1" columns — they share the same month as 착공+1
    # If project_start is known, adjust pre-start columns to the previous month
    if project_start:
        month_group_val = str(ws.cell(month_group_row, schedule_start_col).value or "")
        if "착공" in month_group_val and "-" in month_group_val:
            # Cols schedule_start_col..schedule_start_col+2 are pre-start
            for c in range(schedule_start_col, schedule_start_col + 3):
                if c in col_date_map:
                    existing = col_date_map[c]
                    # Shift to previous month
                    if existing.month == 1:
                        col_date_map[c] = date(existing.year - 1, 12, existing.day)
                    else:
                        col_date_map[c] = date(existing.year, existing.month - 1, min(existing.day, 28))

    return col_date_map


# ---------------------------------------------------------------------------
# Internal: activity scanning
# ---------------------------------------------------------------------------

def _scan_activities(
    ws: Any,
    *,
    col_date_map: dict[int, date],
    first_data_row: int,
    schedule_start_col: int,
    warnings: list[str],
) -> list[ParsedActivity]:
    """Scan rows for activities with bar-chart spans."""
    max_row = ws.max_row or 200
    max_col = ws.max_column or 200

    activities: list[ParsedActivity] = []
    current_discipline = ""

    for r in range(first_data_row, max_row + 1):
        col_a = str(ws.cell(r, 1).value or "").strip().replace("\n", " ")
        col_b = str(ws.cell(r, 2).value or "").strip().replace("\n", " ")
        col_c = str(ws.cell(r, 3).value or "").strip().replace("\n", " ")

        # Update current discipline from column A
        if col_a and col_a not in ("주요 Milestone", "외주 및 자재 발주", "월별 / 누계 공정율"):
            current_discipline = col_a

        # Skip header/summary rows
        if col_a in ("주요 Milestone", "외주 및 자재 발주", "월별 / 누계 공정율"):
            continue

        # An activity row should have content in col B or col C, or a bar in the schedule area
        if not col_b and not col_c:
            # Check if there's a bar anyway (some rows only have bars)
            has_bar = False
            for c in range(schedule_start_col, min(max_col + 1, 180)):
                if ws.cell(r, c).value is not None:
                    has_bar = True
                    break
            if not has_bar:
                continue

        # Scan for bar span
        first_col: int | None = None
        last_col: int | None = None
        bar_texts: list[str] = []

        for c in range(schedule_start_col, min(max_col + 1, 180)):
            val = ws.cell(r, c).value
            if val is not None and str(val).strip():
                if first_col is None:
                    first_col = c
                last_col = c
                bar_texts.append(str(val).strip().replace("\n", " "))

        # Also check merged cells spanning this row
        for merge_range in ws.merged_cells.ranges:
            if merge_range.min_row <= r <= merge_range.max_row and merge_range.min_col >= schedule_start_col:
                if first_col is None or merge_range.min_col < first_col:
                    first_col = merge_range.min_col
                if last_col is None or merge_range.max_col > last_col:
                    last_col = merge_range.max_col

        if first_col is None or last_col is None:
            continue

        # Convert columns to dates
        es_date = _col_to_date(first_col, col_date_map)
        ef_date = _col_to_date(last_col, col_date_map)

        duration_days = 0
        if es_date and ef_date:
            duration_days = max((ef_date - es_date).days, 1)

        item_name = col_b or col_c or ""
        detail = col_c if col_b and col_c and col_b != col_c else ""

        activities.append(
            ParsedActivity(
                row_number=r,
                discipline=current_discipline,
                item=item_name,
                detail=detail,
                es_date=es_date,
                ef_date=ef_date,
                duration_days=duration_days,
                bar_texts=bar_texts,
            )
        )

    return activities


def _merge_continuation_rows(activities: list[ParsedActivity]) -> list[ParsedActivity]:
    """Merge rows without an item name into the preceding named activity.

    In bar-chart schedules, a single activity can span multiple Excel rows.
    Rows with an empty item/detail are continuation rows whose bar extends
    the previous activity's date range.
    """
    merged: list[ParsedActivity] = []
    for a in activities:
        if not a.item and not a.detail and merged:
            # Continuation row → extend previous activity's date range
            prev = merged[-1]
            if a.es_date and (prev.es_date is None or a.es_date < prev.es_date):
                prev.es_date = a.es_date
            if a.ef_date and (prev.ef_date is None or a.ef_date > prev.ef_date):
                prev.ef_date = a.ef_date
            prev.bar_texts.extend(a.bar_texts)
            if prev.es_date and prev.ef_date:
                prev.duration_days = max((prev.ef_date - prev.es_date).days, 1)
        else:
            merged.append(a)
    return merged


def _col_to_date(col: int, col_date_map: dict[int, date]) -> date | None:
    """Find the date for a column, using nearest known column if exact match missing."""
    if col in col_date_map:
        return col_date_map[col]
    # Try nearest column within ±2
    for offset in (1, -1, 2, -2):
        if (col + offset) in col_date_map:
            return col_date_map[col + offset]
    return None


# ---------------------------------------------------------------------------
# Internal: discipline normalization
# ---------------------------------------------------------------------------

_SCHEDULE_DISCIPLINE_MAP: dict[str, str] = {
    "공통가설": "건축",
    "토공사": "토목",
    "타워크레인": "건축",
    "타워크레인 호이스트": "건축",
    "골조공사": "건축",
    "외부석공사": "건축",
    "방수공사": "건축",
    "지하주차장": "건축",
    "세대마감공사": "건축",
    "공용부위 마감공사": "건축",
    "공용부위마감공사": "건축",
    "부대토목": "토목",
    "조경공사": "토목",
    "기계설비": "기계설비",
    "전기공사": "전기설비",
    "소방공사": "소방설비",
    "소방설비": "소방설비",
}


def _normalize_discipline(raw: str) -> str:
    """Map schedule discipline names to smart-scheduler standard disciplines."""
    cleaned = raw.replace("\n", " ").strip()
    if cleaned in _SCHEDULE_DISCIPLINE_MAP:
        return _SCHEDULE_DISCIPLINE_MAP[cleaned]
    # Try partial match (only if cleaned is non-empty to avoid "" in "any" = True)
    if cleaned:
        for key, value in _SCHEDULE_DISCIPLINE_MAP.items():
            if key in cleaned or cleaned in key:
                return value
    return cleaned or "공통"
