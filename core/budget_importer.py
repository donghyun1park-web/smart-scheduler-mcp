"""Import construction budget (실행예산내역서) Excel into smart-scheduler DB.

Supports two common formats:
- **Type A** (PC공사, 가설/토목): 20 cols, col A=sortMngNo, col B=명칭, col F=구분(당초/변경/증감)
- **Type B** (기계설비, 전기/통신/소방): 27 cols, col A=명칭, col H=구분(당초/변경/증감)

Each item appears in 3 consecutive rows: 당초 (original), 변경 (revised), 증감 (delta).
The importer extracts the "변경" row as the current budget.

Items are grouped into cost categories (공종) by detecting hierarchy rows
(rows where 관리번호 is '-' or empty = summary; rows with a real code = detail item).
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from core.number_utils import to_float


# ---------------------------------------------------------------------------
# Public dataclasses
# ---------------------------------------------------------------------------

@dataclass
class ParsedBudgetItem:
    """A single budget line item extracted from 실행예산내역서."""

    row_number: int
    name: str
    spec: str
    unit: str
    management_code: str
    category: str
    gubun: str  # 당초/변경/증감
    contract_qty: float
    contract_unit_price: float
    contract_amount: float
    execution_qty: float
    execution_unit_price: float
    execution_amount: float
    cost_type: str  # 비목 (외주비, 자재비, etc.)
    discipline: str


@dataclass
class BudgetCategory:
    """Aggregated budget by category (공종 level)."""

    category: str
    discipline: str
    contract_amount: float
    execution_amount: float
    item_count: int


@dataclass
class BudgetParseResult:
    """Result of parsing a budget workbook."""

    file_name: str
    discipline: str
    items: list[ParsedBudgetItem]
    categories: list[BudgetCategory]
    total_contract_amount: float
    total_execution_amount: float
    warnings: list[str]


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------

def parse_budget_excel(
    path: str | Path,
    *,
    discipline: str = "",
    sheet_name: str | None = None,
) -> BudgetParseResult:
    """Parse a 실행예산내역서 Excel and return structured budget items.

    Parameters
    ----------
    path : file path to the .xlsx workbook
    discipline : override discipline name (e.g., '기계설비', '전기설비')
    sheet_name : which sheet to read (default: first sheet)
    """
    path = Path(path)
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    warnings: list[str] = []

    # Detect format type from header row
    format_type = _detect_format(ws)
    if not discipline:
        discipline = _infer_discipline(path.name)

    items = _parse_items(ws, format_type=format_type, discipline=discipline, warnings=warnings)
    categories = _aggregate_categories(items)

    total_contract = sum(cat.contract_amount for cat in categories)
    total_execution = sum(cat.execution_amount for cat in categories)

    wb.close()
    return BudgetParseResult(
        file_name=path.name,
        discipline=discipline,
        items=items,
        categories=categories,
        total_contract_amount=total_contract,
        total_execution_amount=total_execution,
        warnings=warnings,
    )


def import_budgets_to_db(
    db_path: str | Path,
    budget_paths: list[str | Path],
    *,
    disciplines: list[str] | None = None,
    project_id: str | None = None,
) -> dict[str, Any]:
    """Parse multiple budget Excel files and write CostItems into a .scheduler DB.

    The budgets are aggregated by category and mapped to existing activities
    by matching discipline. If no matching activity exists, a warning is issued.

    Returns a summary dict.
    """
    from core import db
    from core.models import CostItem

    all_categories: list[BudgetCategory] = []
    all_warnings: list[str] = []

    for idx, bp in enumerate(budget_paths):
        disc = disciplines[idx] if disciplines and idx < len(disciplines) else ""
        result = parse_budget_excel(bp, discipline=disc)
        all_categories.extend(result.categories)
        all_warnings.extend(result.warnings)

    # Try to match categories to existing activities by discipline
    activities = db.list_activities(db_path)
    activity_by_discipline: dict[str, list[Any]] = {}
    for act in activities:
        activity_by_discipline.setdefault(act.discipline, []).append(act)

    cost_items_created = 0
    unmatched: list[str] = []

    for cat in all_categories:
        norm_disc = _normalize_budget_discipline(cat.discipline)
        matched_activities = activity_by_discipline.get(norm_disc, [])

        if not matched_activities:
            # Try broader match
            for key, acts in activity_by_discipline.items():
                if norm_disc in key or key in norm_disc:
                    matched_activities = acts
                    break

        if matched_activities:
            # Distribute budget evenly across matching activities
            share = len(matched_activities)
            for act in matched_activities:
                cost_item_id = f"cost-{act.activity_id}"
                existing = db.list_cost_items(db_path, activity_id=act.activity_id)
                contract = cat.contract_amount / share
                execution = cat.execution_amount / share

                if existing:
                    # Add to existing
                    ex = existing[0]
                    db.upsert_cost_item(
                        db_path,
                        CostItem(
                            cost_item_id=ex.cost_item_id,
                            activity_id=act.activity_id,
                            contract_amount=ex.contract_amount + contract,
                            execution_budget=ex.execution_budget + execution,
                            invested_cost=ex.invested_cost,
                            billing_amount=ex.billing_amount,
                        ),
                    )
                else:
                    db.upsert_cost_item(
                        db_path,
                        CostItem(
                            cost_item_id=cost_item_id,
                            activity_id=act.activity_id,
                            contract_amount=contract,
                            execution_budget=execution,
                        ),
                    )
                cost_items_created += 1
        else:
            unmatched.append(f"{cat.discipline}/{cat.category}: contract={cat.contract_amount:,.0f}, execution={cat.execution_amount:,.0f}")

    if unmatched:
        all_warnings.append(f"Unmatched categories (no activity with matching discipline): {len(unmatched)}")
        for item in unmatched[:10]:
            all_warnings.append(f"  - {item}")

    return {
        "ok": True,
        "cost_items_created": cost_items_created,
        "categories_total": len(all_categories),
        "unmatched_categories": len(unmatched),
        "total_contract": sum(c.contract_amount for c in all_categories),
        "total_execution": sum(c.execution_amount for c in all_categories),
        "warnings": all_warnings,
        "db_path": str(db_path),
    }


# ---------------------------------------------------------------------------
# Internal: format detection
# ---------------------------------------------------------------------------

_FORMAT_A_MARKERS = {"sortMngNo", "sortmngno"}
_FORMAT_B_MARKERS = {"자원구분", "지급구분", "간접비여부"}


def _detect_format(ws: Any) -> str:
    """Detect whether this is Type A (20-col) or Type B (27-col) format."""
    row1_vals = {str(ws.cell(1, c).value or "").strip().lower() for c in range(1, 10)}
    if row1_vals & _FORMAT_A_MARKERS:
        return "A"
    row1_vals_full = {str(ws.cell(1, c).value or "").strip() for c in range(1, 10)}
    if row1_vals_full & _FORMAT_B_MARKERS:
        return "B"
    # Heuristic: check column count
    if (ws.max_column or 0) > 22:
        return "B"
    return "A"


def _infer_discipline(filename: str) -> str:
    """Infer discipline from the filename."""
    name = filename.lower()
    if "기계" in name or "설비" in name:
        return "기계설비"
    if "전기" in name:
        return "전기설비"
    if "소방" in name:
        return "소방설비"
    if "토목" in name or "가설" in name:
        return "토목"
    if "건축" in name or "pc" in name.lower():
        return "건축"
    return "공통"


# ---------------------------------------------------------------------------
# Internal: item parsing
# ---------------------------------------------------------------------------

def _parse_items(
    ws: Any,
    *,
    format_type: str,
    discipline: str,
    warnings: list[str],
) -> list[ParsedBudgetItem]:
    """Parse '변경' rows, carrying forward name/spec/unit from the preceding '당초' row."""
    items: list[ParsedBudgetItem] = []
    max_row = ws.max_row or 1
    current_category = ""

    # State: carry forward from the most recent '당초' row
    pending_name = ""
    pending_spec = ""
    pending_unit = ""
    pending_mgmt = ""
    pending_cost_type = ""

    for r in range(3, max_row + 1):
        row_data = _read_row(ws, r, format_type)
        if row_data is None:
            continue

        name = row_data["name"]
        mgmt_code = row_data["management_code"]
        gubun = row_data["gubun"]

        # '당초' row: remember the item identity for the upcoming '변경' row
        if gubun == "당초":
            if name:
                pending_name = name
                pending_spec = row_data["spec"]
                pending_unit = row_data["unit"]
                pending_mgmt = mgmt_code
                pending_cost_type = row_data.get("cost_type", "")
            # Track category from summary rows (mgmt_code is '-' or empty)
            if mgmt_code in ("-", "") and name and name not in ("직접비",):
                current_category = name
            continue

        # Only take '변경' rows
        if gubun != "변경":
            continue

        # Use carried-forward identity if this row's name is empty
        effective_name = name or pending_name
        effective_spec = row_data["spec"] or pending_spec
        effective_unit = row_data["unit"] or pending_unit
        effective_mgmt = mgmt_code or pending_mgmt
        effective_cost_type = row_data.get("cost_type", "") or pending_cost_type

        contract_amount = to_float(row_data["contract_amount"])
        execution_amount = to_float(row_data["execution_amount"])

        # Skip rows with zero amounts and no management code (empty summaries)
        if contract_amount == 0 and execution_amount == 0 and effective_mgmt in ("-", ""):
            continue

        items.append(
            ParsedBudgetItem(
                row_number=r,
                name=effective_name,
                spec=effective_spec,
                unit=effective_unit,
                management_code=effective_mgmt,
                category=current_category,
                gubun=gubun,
                contract_qty=to_float(row_data["contract_qty"]),
                contract_unit_price=to_float(row_data["contract_unit_price"]),
                contract_amount=contract_amount,
                execution_qty=to_float(row_data["execution_qty"]),
                execution_unit_price=to_float(row_data["execution_unit_price"]),
                execution_amount=execution_amount,
                cost_type=effective_cost_type,
                discipline=discipline,
            )
        )

    return items


def _read_row(ws: Any, r: int, format_type: str) -> dict[str, Any] | None:
    """Read a single row and return normalized field dict, or None if completely empty.

    Note: '변경' and '증감' rows typically have an empty name; the name is
    carried forward from the preceding '당초' row by the caller.
    """
    if format_type == "A":
        # Type A: sortMngNo(A), 명칭(B), 규격(C), 단위(D), 관리번호(E), 구분(F),
        #         도급수량(G), 도급단가(H), 도급금액(I), 실행수량(J), 실행단가(K), 실행금액(L),
        #         실행율(M), 비목(N), 자재비(O)
        name = str(ws.cell(r, 2).value or "").strip()
        gubun = str(ws.cell(r, 6).value or "").strip()
        # A row is valid if it has a name OR a gubun (변경/증감 rows lack name)
        if not name and not gubun:
            return None
        return {
            "name": name,
            "spec": str(ws.cell(r, 3).value or "").strip(),
            "unit": str(ws.cell(r, 4).value or "").strip(),
            "management_code": str(ws.cell(r, 5).value or "").strip(),
            "gubun": gubun,
            "contract_qty": ws.cell(r, 7).value,
            "contract_unit_price": ws.cell(r, 8).value,
            "contract_amount": ws.cell(r, 9).value,
            "execution_qty": ws.cell(r, 10).value,
            "execution_unit_price": ws.cell(r, 11).value,
            "execution_amount": ws.cell(r, 12).value,
            "cost_type": str(ws.cell(r, 14).value or "").strip(),
        }
    else:
        # Type B: 명칭(A), 규격(B), 단위(C), 관리번호(D), 자원구분(E), 지급구분(F),
        #         간접비여부(G), 구분(H), 승인번호(I),
        #         도급수량(J), 도급단가(K), 도급금액(L), 실행수량(M), 실행단가(N), 실행금액(O),
        #         하도금액(P), ...
        name = str(ws.cell(r, 1).value or "").strip()
        gubun = str(ws.cell(r, 8).value or "").strip()
        if not name and not gubun:
            return None
        return {
            "name": name,
            "spec": str(ws.cell(r, 2).value or "").strip(),
            "unit": str(ws.cell(r, 3).value or "").strip(),
            "management_code": str(ws.cell(r, 4).value or "").strip(),
            "gubun": gubun,
            "contract_qty": ws.cell(r, 10).value,
            "contract_unit_price": ws.cell(r, 11).value,
            "contract_amount": ws.cell(r, 12).value,
            "execution_qty": ws.cell(r, 13).value,
            "execution_unit_price": ws.cell(r, 14).value,
            "execution_amount": ws.cell(r, 15).value,
            "cost_type": str(ws.cell(r, 5).value or "").strip(),
        }


# ---------------------------------------------------------------------------
# Internal: aggregation
# ---------------------------------------------------------------------------

def _aggregate_categories(items: list[ParsedBudgetItem]) -> list[BudgetCategory]:
    """Aggregate items by category into BudgetCategory summaries."""
    groups: dict[str, dict[str, Any]] = {}
    for item in items:
        key = item.category or item.name
        if key not in groups:
            groups[key] = {
                "category": key,
                "discipline": item.discipline,
                "contract_amount": 0.0,
                "execution_amount": 0.0,
                "item_count": 0,
            }
        groups[key]["contract_amount"] += item.contract_amount
        groups[key]["execution_amount"] += item.execution_amount
        groups[key]["item_count"] += 1

    return [
        BudgetCategory(
            category=g["category"],
            discipline=g["discipline"],
            contract_amount=round(g["contract_amount"], 2),
            execution_amount=round(g["execution_amount"], 2),
            item_count=g["item_count"],
        )
        for g in groups.values()
    ]


# ---------------------------------------------------------------------------
# Internal: discipline normalization
# ---------------------------------------------------------------------------

_BUDGET_DISCIPLINE_MAP: dict[str, str] = {
    "기계설비": "기계설비",
    "기계": "기계설비",
    "설비": "기계설비",
    "전기설비": "전기설비",
    "전기": "전기설비",
    "소방설비": "소방설비",
    "소방": "소방설비",
    "토목": "토목",
    "건축": "건축",
    "가설": "건축",
    "공통": "공통",
}


def _normalize_budget_discipline(raw: str) -> str:
    """Normalize budget discipline to match scheduler standard."""
    cleaned = raw.strip()
    if cleaned in _BUDGET_DISCIPLINE_MAP:
        return _BUDGET_DISCIPLINE_MAP[cleaned]
    for key, value in _BUDGET_DISCIPLINE_MAP.items():
        if key in cleaned:
            return value
    return cleaned or "공통"
