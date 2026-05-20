from __future__ import annotations

from datetime import date

from openpyxl import load_workbook

from core import db
from core.models import Activity, Calendar, Project, Relationship, WBS
from tools.analysis_tools import calculate_cpm
from tools.report_tools import generate_report


REPORT_SHEETS = [
    "\uc694\uc57d",
    "WBS",
    "Activity \uc0c1\uc138",
    "\uad00\uacc4",
    "S-Curve \ub370\uc774\ud130",
]


def test_generate_report_creates_five_sheet_excel_with_critical_highlight(tmp_path):
    db_path = _project_db(tmp_path)
    calculate_cpm(db_path)

    result = generate_report(db_path, tmp_path / "report.xlsx")

    assert result["ok"] is True
    assert result["sheets"] == REPORT_SHEETS
    workbook = load_workbook(result["output_path"])
    assert workbook.sheetnames == result["sheets"]
    activity_sheet = workbook["Activity \uc0c1\uc138"]
    critical_fill = activity_sheet["A2"].fill.fgColor.rgb
    assert critical_fill in {"FFFFCCCC", "00FFCCCC"}
    workbook.close()


def test_end_to_end_project_to_cpm_gantt_scurve_report(tmp_path):
    from core.s_curve import build_s_curve_data
    from core.visualization import build_gantt_figure

    db_path = _project_db(tmp_path)
    calculate_cpm(db_path)
    activities = db.list_activities(db_path)

    gantt = build_gantt_figure(activities)
    s_curve = build_s_curve_data(activities)
    report = generate_report(db_path)

    assert len(gantt.data) == 2
    assert s_curve[-1]["cumulative_value"] == 300.0
    assert report["ok"] is True


def _project_db(tmp_path):
    db_path = tmp_path / "report.scheduler"
    db.initialize_database(db_path)
    db.create_calendar(db_path, Calendar("cal", "Korean 5-day", "1111100"))
    db.create_project(db_path, Project("project", "Report Pilot", date(2026, 1, 5), "cal"))
    db.add_wbs(db_path, WBS("wbs", None, "MEP", "MEP", 1))
    db.add_activity(db_path, _activity("a", "A100", "Sleeve", 1, 100.0))
    db.add_activity(db_path, _activity("b", "A200", "Duct", 2, 200.0))
    db.add_relationship(db_path, Relationship("r1", "a", "b", "FS", 0))
    return db_path


def _activity(activity_id: str, code: str, name: str, duration: int, cost: float) -> Activity:
    return Activity(
        activity_id=activity_id,
        code=code,
        name=name,
        wbs_id="wbs",
        discipline="\uacf5\ud1b5",
        zone="1F",
        duration=duration,
        cost=cost,
    )
